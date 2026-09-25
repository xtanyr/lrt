from pathlib import Path
import sys
from openpyxl import load_workbook

SOURCE = Path(r"C:\Users\-EAFC~1\AppData\Local\Temp\browser-use\exports")
sys.stdout.reconfigure(encoding="utf-8")

filter_text = sys.argv[1].casefold() if len(sys.argv) > 1 else ""
for path in sorted(SOURCE.glob("*.xlsx")):
    if filter_text and filter_text not in path.name.casefold():
        continue
    print(f"WORKBOOK {path.name}")
    wb = load_workbook(path, read_only=False, data_only=False)
    for ws in wb.worksheets:
        print(f"SHEET {ws.title!r} state={ws.sheet_state} size={ws.max_row}x{ws.max_column}")
        print("MERGES", ",".join(str(r) for r in ws.merged_cells.ranges))
        print("HIDDEN_ROWS", [i for i, d in ws.row_dimensions.items() if d.hidden][:30])
        print("HIDDEN_COLS", [i for i, d in ws.column_dimensions.items() if d.hidden][:30])
        if ws.title.casefold() == "метрики":
            row_numbers = list(range(1, 13)) + list(range(340, 361))
            max_column = min(ws.max_column, 17)
        elif ws.title.casefold() != "рейтинг":
            row_numbers = list(range(1, min(ws.max_row, 80) + 1))
            max_column = min(ws.max_column, 12)
        else:
            row_numbers = list(range(1, min(ws.max_row, 30) + 1))
            max_column = min(ws.max_column, 20)
        for row_number in row_numbers:
            cells = []
            for row in ws.iter_rows(min_row=row_number, max_row=row_number, max_col=max_column):
             for cell in row:
                value = cell.value
                if value is None:
                    continue
                if isinstance(value, str):
                    display = value.replace("\n", "\\n")
                else:
                    display = "<NUMBER>"
                if isinstance(value, str) and value.startswith("="):
                    display = "<FORMULA>"
                cells.append(f"{cell.coordinate}={display!r}[fmt={cell.number_format!r}]")
            if cells:
                print("ROW", " | ".join(cells))
        print()
    print("END_WORKBOOK")
